import calendar
import json
import re
import sys
from pathlib import Path

import openpyxl

MONTH_NAMES = [
    None,
    "january", "february", "march", "april", "may", "june",
    "july", "august", "september", "october", "november", "december",
]

WEEKDAY_PATTERN = re.compile(
    r"^(Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\b.*\(([^)]*)\)",
    re.IGNORECASE,
)

MEAL_COLUMNS = {
    2: "breakfast",
    3: "lunch",
    5: "snacks",
    6: "dinner",
}


def clean_cell(value):
    if value is None:
        return []

    if isinstance(value, (int, float)):
        if value == 0:
            return []
        value = str(value)

    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return []

    # Some women's-menu cells contain multiple menu items separated by
    # line breaks. Treat each line as a separate item.
    items = []
    for part in text.split("\n"):
        item = re.sub(r"\s+", " ", part).strip()
        if item and item != "0":
            items.append(item)

    return items


def parse_date_block(value, year, month):
    if not isinstance(value, str):
        return None

    match = WEEKDAY_PATTERN.search(value.replace("\n", " "))
    if not match:
        return None

    weekday = match.group(1).capitalize()
    max_day = calendar.monthrange(year, month)[1]

    days = [
        int(number)
        for number in re.findall(r"\d+", match.group(2))
        if 1 <= int(number) <= max_day
    ]

    if not days:
        return None

    return weekday, days


def find_date_blocks(ws, year, month):
    blocks = []

    for row in range(4, ws.max_row + 1):
        parsed = parse_date_block(ws.cell(row, 1).value, year, month)
        if parsed:
            weekday, days = parsed
            blocks.append((row, weekday, days))

    return blocks


def parse_sheet(ws, year, month):
    blocks = find_date_blocks(ws, year, month)
    menus = {}

    for index, (start_row, weekday, days) in enumerate(blocks):
        # The next valid date label starts the next menu block.
        # The menu itself normally ends before the instructions section.
        if index + 1 < len(blocks):
            end_row = blocks[index + 1][0] - 1
        else:
            end_row = min(171, ws.max_row)

        meals = {
            "breakfast": [],
            "lunch": [],
            "snacks": [],
            "dinner": [],
        }

        for row in range(start_row, end_row + 1):
            for column, meal in MEAL_COLUMNS.items():
                meals[meal].extend(clean_cell(ws.cell(row, column).value))

        for day in days:
            date_string = f"{year}-{month:02d}-{day:02d}"
            menus[date_string] = {
                "day": weekday[:3],
                "meals": {
                    meal: list(items)
                    for meal, items in meals.items()
                },
            }

    return dict(sorted(menus.items()))


def convert_womens_workbook(input_file):
    input_path = Path(input_file)

    match = re.match(
        r"^(january|february|march|april|may|june|july|august|"
        r"september|october|november|december)_womens_(\d{4})\.xlsx$",
        input_path.name,
        re.IGNORECASE,
    )

    if not match:
        raise ValueError(
            "Expected filename like october_womens_2026.xlsx"
        )

    month_name = match.group(1).lower()
    year = int(match.group(2))
    month = MONTH_NAMES.index(month_name)

    workbook = openpyxl.load_workbook(input_path, data_only=True)

    required_sheets = ["Veg & Non Veg", "Spl Mess"]
    for sheet in required_sheets:
        if sheet not in workbook.sheetnames:
            raise ValueError(
                f"Missing required sheet '{sheet}'. "
                f"Found: {workbook.sheetnames}"
            )

    output_dir = (
        Path("output")
        / "womens"
        / str(year)
        / f"{month:02d}"
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    sheet_config = [
        ("Veg & Non Veg", "veg_non_veg", "veg-non-veg"),
        ("Spl Mess", "special", "special"),
    ]

    for sheet_name, mess_type, suffix in sheet_config:
        menus = parse_sheet(
            workbook[sheet_name],
            year,
            month,
        )

        data = {
            "hostel": "womens",
            "year": year,
            "month": month,
            "month_name": month_name.capitalize(),
            "mess_type": mess_type,
            "menus": menus,
        }

        output_file = (
            output_dir
            / f"{month_name}_womens_{year}_{suffix}.json"
        )

        output_file.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        print(
            f"Created {output_file} "
            f"({len(menus)} dates)"
        )


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(
            "Usage: python convert_womens_menu.py "
            "october_womens_2026.xlsx"
        )
        sys.exit(1)

    convert_womens_workbook(sys.argv[1])
